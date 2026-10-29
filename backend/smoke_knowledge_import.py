"""教材提取回归：真实 PDF 结构与 OCR 协议替身，不访问模型或业务数据库。"""
import hashlib
import io
import json
import subprocess
import tempfile
import threading
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pymupdf

from app.services import knowledge_import as importer
from app.services import knowledge_ocr as ocr


class KnowledgeExtractionTests(unittest.TestCase):
    def setUp(self):
        self.enterContext(patch.multiple(
            importer.settings, RAG_OCR_DPI=96, RAG_OCR_WORKERS=2, RAG_VISUAL_ENABLED=True,
        ))
        self.recognize = self.enterContext(patch.object(
            importer, "recognize_text", side_effect=AssertionError("unexpected OCR call"),
        ))
        self.describe = self.enterContext(patch.object(
            importer, "describe_visual", side_effect=AssertionError("unexpected visual request"),
        ))

    @staticmethod
    def image_bytes(text="Scanned equation and solution."):
        with pymupdf.open() as source:
            page = source.new_page(width=240, height=300)
            page.insert_textbox(pymupdf.Rect(20, 60, 220, 200), text, fontsize=12)
            return page.get_pixmap(alpha=False).tobytes("png")

    @classmethod
    def pdf_bytes(cls, *kinds, scan_text="Scanned equation and solution.", encrypted=False):
        image = cls.image_bytes(scan_text)
        with pymupdf.open() as document:
            for kind in kinds:
                page = document.new_page(width=240, height=300)
                if kind in {"native", "mixed", "header_scan", "vector"}:
                    page.insert_text((20, 30), "Native lesson heading.", fontsize=10)
                if kind in {"scan", "header_scan"}:
                    page.insert_image(page.rect, stream=image)
                elif kind == "mixed":
                    page.insert_image(pymupdf.Rect(20, 90, 200, 260), stream=image)
                elif kind == "vector":
                    page.draw_polyline([(30, 200), (190, 200), (30, 110), (30, 200)])
            options = {"encryption": pymupdf.PDF_ENCRYPT_AES_256,
                       "owner_pw": "owner-test", "user_pw": "reader-test"} if encrypted else {}
            return document.tobytes(**options)

    def mock_ocr(self, text="方程移项时必须改变符号。", confidence=0.95):
        self.recognize.side_effect = None
        self.recognize.return_value = {"text": text, "confidence": confidence}

    @classmethod
    def multi_image_pdf(cls):
        with pymupdf.open() as document:
            page = document.new_page(width=240, height=300)
            page.insert_text((20, 30), "Native lesson heading.", fontsize=10)
            page.insert_image(pymupdf.Rect(20, 60, 110, 160), stream=cls.image_bytes("First diagram."))
            page.insert_image(pymupdf.Rect(130, 180, 220, 280), stream=cls.image_bytes("Second diagram."))
            return document.tobytes()

    def test_native_text_bypasses_ocr_and_visual_models(self):
        updates = []
        result = importer.extract_knowledge_pdf(self.pdf_bytes("native"), progress=updates.append)
        self.assertIn("Native lesson heading.", result["pages"][0][1])
        self.assertEqual((result["page_count"], result["indexed_pages"], result["ocr_pages"], result["visual_pages"]),
                         (1, 1, 0, 0))
        self.recognize.assert_not_called()
        self.describe.assert_not_called()
        self.assertEqual(updates[-1]["completed_pages"], 1)
        self.assertEqual(updates[-1]["total_pages"], 1)

    def test_mixed_page_preserves_native_text_and_ocr_image_words(self):
        self.mock_ocr("图内文字需要补充。")
        self.describe.side_effect = None
        self.describe.return_value = "图中两条直线互相垂直。"
        result = importer.extract_knowledge_pdf(self.pdf_bytes("mixed"))
        text = result["pages"][0][1]
        self.assertIn("Native lesson heading.", text)
        self.assertIn("图内文字需要补充。", text)
        self.assertIn("图中两条直线互相垂直。", text)
        self.recognize.assert_called_once()
        self.describe.assert_called_once()
        self.assertIn("图内文字需要补充。", self.describe.call_args.args[1])

    def test_native_header_does_not_hide_full_page_scanned_body(self):
        self.mock_ocr("扫描正文包含一元二次方程的解法。")
        with patch.object(importer.settings, "RAG_VISUAL_ENABLED", False):
            result = importer.extract_knowledge_pdf(self.pdf_bytes("header_scan"))
        self.assertIn("Native lesson heading.", result["pages"][0][1])
        self.assertIn("扫描正文包含一元二次方程的解法。", result["pages"][0][1])
        self.assertEqual(result["ocr_pages"], 1)
        self.recognize.assert_called_once()

    def test_successful_scan_checkpoint_skips_all_repeated_ocr(self):
        self.mock_ocr()
        data = self.pdf_bytes("scan", "scan")
        with tempfile.TemporaryDirectory() as directory:
            first = importer.extract_knowledge_pdf(data, checkpoint_dir=directory)
            self.assertEqual(self.recognize.call_count, 2)
            self.recognize.reset_mock()
            self.recognize.side_effect = AssertionError("completed pages must be reused")
            second = importer.extract_knowledge_pdf(data, checkpoint_dir=directory)
            self.assertEqual(first, second)
            self.recognize.assert_not_called()
            self.assertEqual(len(list(Path(directory).glob("*/page-*.json"))), 2)
        self.describe.assert_not_called()

    def test_checkpoint_directory_is_scoped_by_file_digest(self):
        first_data = self.pdf_bytes("scan", scan_text="First source content.")
        second_data = self.pdf_bytes("scan", scan_text="Second source content.")
        self.assertNotEqual(hashlib.sha256(first_data).hexdigest(), hashlib.sha256(second_data).hexdigest())
        self.mock_ocr("第一份文档文字。")
        with tempfile.TemporaryDirectory() as directory:
            importer.extract_knowledge_pdf(first_data, checkpoint_dir=directory)
            self.mock_ocr("第二份文档文字。")
            second = importer.extract_knowledge_pdf(second_data, checkpoint_dir=directory)
            self.assertIn("第二份文档文字。", second["pages"][0][1])
            self.assertNotIn("第一份文档文字。", second["pages"][0][1])
            self.assertEqual(self.recognize.call_count, 2)
            self.assertEqual(len(list(Path(directory).iterdir())), 2)

    def test_retry_reuses_successful_pages_and_retries_failed_page_only(self):
        data = self.pdf_bytes("scan", "scan", "scan")
        # One worker fixes page order while exercising the production executor.
        self.recognize.side_effect = [
            {"text": "第一页成功。", "confidence": 0.95},
            RuntimeError("页面文字识别超时。"),
            {"text": "第三页成功。", "confidence": 0.95},
        ]
        with tempfile.TemporaryDirectory() as directory, patch.object(importer.settings, "RAG_OCR_WORKERS", 1):
            with self.assertRaisesRegex(RuntimeError, "第 2 页文字识别失败"):
                importer.extract_knowledge_pdf(data, checkpoint_dir=directory)
            self.assertEqual(self.recognize.call_count, 3)
            self.recognize.reset_mock()
            self.mock_ocr("第二页重试成功。")
            result = importer.extract_knowledge_pdf(data, checkpoint_dir=directory)
        self.recognize.assert_called_once()
        self.assertEqual([number for number, _ in result["pages"]], [1, 2, 3])
        self.assertIn("第一页成功。", result["pages"][0][1])
        self.assertIn("第二页重试成功。", result["pages"][1][1])
        self.assertIn("第三页成功。", result["pages"][2][1])

    def test_scan_diagram_uses_dedicated_ocr_before_visual_semantics(self):
        self.mock_ocr("如图，三角形 ABC 是直角三角形。")
        self.describe.side_effect = None
        self.describe.return_value = "AB 与 BC 垂直，AC 为斜边。"
        result = importer.extract_knowledge_pdf(self.pdf_bytes("scan"))
        self.recognize.assert_called_once()
        self.describe.assert_called_once()
        self.assertIn("三角形 ABC", self.describe.call_args.args[1])
        self.assertIn("三角形 ABC", result["pages"][0][1])
        self.assertIn("AC 为斜边", result["pages"][0][1])
        self.assertEqual((result["ocr_pages"], result["visual_pages"]), (1, 1))

    def test_visual_retry_reuses_successful_ocr_checkpoint(self):
        self.mock_ocr("如图，三角形 ABC 是直角三角形。")
        self.describe.side_effect = RuntimeError("图形服务不可用。")
        data = self.pdf_bytes("scan")
        updates = []
        with tempfile.TemporaryDirectory() as directory:
            first = importer.extract_knowledge_pdf(data, checkpoint_dir=directory, progress=updates.append)
            checkpoint = json.loads(next(Path(directory).glob("*/page-1.json")).read_text(encoding="utf-8"))
            self.assertFalse(checkpoint["complete"])
            self.assertTrue(checkpoint["ocr_done"])
            self.assertIn("三角形 ABC", checkpoint["text"])
            self.assertIn("三角形 ABC", first["pages"][0][1])
            self.assertTrue(any("图形解析暂不可用" in warning for warning in first["warnings"]))
            self.assertEqual(updates[-1]["completed_pages"], 1)
            self.recognize.reset_mock()
            self.recognize.side_effect = AssertionError("retry must keep OCR")
            self.describe.side_effect = None
            self.describe.return_value = "AB 垂直于 BC。"
            second = importer.extract_knowledge_pdf(data, checkpoint_dir=directory)
        self.recognize.assert_not_called()
        self.assertEqual(second["visual_pages"], 1)
        self.assertIn("AB 垂直于 BC", second["pages"][0][1])
        self.assertFalse(any("图形解析暂不可用" in warning for warning in second["warnings"]))

    def test_mixed_page_ocr_failure_remains_retriable_despite_native_heading(self):
        with patch.object(importer.settings, "RAG_VISUAL_ENABLED", False):
            for kind in ("mixed", "header_scan"):
                with self.subTest(kind=kind), tempfile.TemporaryDirectory() as directory:
                    self.recognize.side_effect = RuntimeError("暂时无法识别扫描正文。")
                    data = self.pdf_bytes(kind)
                    with self.assertRaisesRegex(RuntimeError, "第 1 页文字识别失败"):
                        importer.extract_knowledge_pdf(data, checkpoint_dir=directory)
                    self.recognize.reset_mock()
                    self.mock_ocr("重试成功的扫描正文。")
                    result = importer.extract_knowledge_pdf(data, checkpoint_dir=directory)
                    self.recognize.assert_called_once()
                    self.assertIn("Native lesson heading.", result["pages"][0][1])
                    self.assertIn("重试成功的扫描正文。", result["pages"][0][1])

    def test_partial_image_ocr_retry_reuses_successful_region_on_same_page(self):
        self.recognize.side_effect = [
            {"text": "第一幅图的文字。", "confidence": 0.95},
            RuntimeError("第二幅图暂时识别失败。"),
        ]
        data = self.multi_image_pdf()
        with tempfile.TemporaryDirectory() as directory, patch.object(importer.settings, "RAG_VISUAL_ENABLED", False):
            with self.assertRaisesRegex(RuntimeError, "第 1 页文字识别失败"):
                importer.extract_knowledge_pdf(data, checkpoint_dir=directory)
            self.assertEqual(self.recognize.call_count, 2)
            checkpoint = json.loads(next(Path(directory).glob("*/page-1.json")).read_text(encoding="utf-8"))
            self.assertFalse(checkpoint.get("ocr_done", False))
            self.assertEqual(len(checkpoint["ocr_parts"]), 1)
            self.recognize.reset_mock()
            self.mock_ocr("第二幅图重试成功的文字。")
            result = importer.extract_knowledge_pdf(data, checkpoint_dir=directory)
        self.recognize.assert_called_once()
        text = result["pages"][0][1]
        self.assertIn("Native lesson heading.", text)
        self.assertIn("第一幅图的文字。", text)
        self.assertIn("第二幅图重试成功的文字。", text)

    def test_partial_visual_retry_reuses_successful_region_and_ocr(self):
        self.mock_ocr("图中注释文字。")
        self.describe.side_effect = ["第一幅图中 AB 垂直于 BC。", RuntimeError("第二幅图解析超时。")]
        data = self.multi_image_pdf()
        with tempfile.TemporaryDirectory() as directory:
            first = importer.extract_knowledge_pdf(data, checkpoint_dir=directory)
            self.assertEqual(self.describe.call_count, 2)
            self.assertEqual(self.recognize.call_count, 2)
            checkpoint = json.loads(next(Path(directory).glob("*/page-1.json")).read_text(encoding="utf-8"))
            self.assertTrue(checkpoint["ocr_done"])
            self.assertFalse(checkpoint["complete"])
            self.assertEqual(len(checkpoint["visual_parts"]), 1)
            self.assertIn("图中注释文字。", first["pages"][0][1])
            self.assertIn("第一幅图中 AB 垂直于 BC。", first["pages"][0][1])
            self.assertTrue(any("图形解析暂不可用" in warning for warning in first["warnings"]))
            self.recognize.reset_mock()
            self.recognize.side_effect = AssertionError("successful OCR must be reused")
            self.describe.reset_mock()
            self.describe.side_effect = None
            self.describe.return_value = "第二幅图中 DE 平行于 FG。"
            result = importer.extract_knowledge_pdf(data, checkpoint_dir=directory)
        self.recognize.assert_not_called()
        self.describe.assert_called_once()
        self.assertIn("第一幅图中 AB 垂直于 BC。", result["pages"][0][1])
        self.assertIn("第二幅图中 DE 平行于 FG。", result["pages"][0][1])

    def test_low_confidence_warning_does_not_invalidate_completed_visual_checkpoint(self):
        self.mock_ocr("方程 x = 2。", confidence=0.5)
        self.describe.side_effect = None
        self.describe.return_value = "等号左侧为未知数 x。"
        data = self.pdf_bytes("scan")
        with tempfile.TemporaryDirectory() as directory:
            first = importer.extract_knowledge_pdf(data, checkpoint_dir=directory)
            self.assertTrue(any("置信度" in warning for warning in first["warnings"]))
            self.recognize.reset_mock()
            self.describe.reset_mock()
            second = importer.extract_knowledge_pdf(data, checkpoint_dir=directory)
        self.assertEqual(first, second)
        self.recognize.assert_not_called()
        self.describe.assert_not_called()

    def test_page_workers_run_concurrently_with_configured_bound(self):
        barrier = threading.Barrier(2, timeout=5)
        lock = threading.Lock()
        counts = {"active": 0, "maximum": 0}

        def recognize(_image):
            with lock:
                counts["active"] += 1
                counts["maximum"] = max(counts["maximum"], counts["active"])
            try:
                barrier.wait()
                return {"text": "可检索扫描文字。", "confidence": 0.95}
            finally:
                with lock:
                    counts["active"] -= 1

        self.recognize.side_effect = recognize
        with patch.object(importer.settings, "RAG_VISUAL_ENABLED", False):
            result = importer.extract_knowledge_pdf(self.pdf_bytes(*(["scan"] * 6)))
        self.assertEqual(counts["maximum"], 2)
        self.assertEqual(self.recognize.call_count, 6)
        self.assertEqual(result["indexed_pages"], 6)

    def test_docx_embedded_png_is_extracted_once_and_keeps_native_text(self):
        from docx import Document

        document = Document()
        document.add_paragraph("原生 Word 教材说明。")
        picture = self.image_bytes("Embedded lesson diagram.")
        document.add_picture(io.BytesIO(picture))
        document.add_picture(io.BytesIO(picture))
        buffer = io.BytesIO()
        document.save(buffer)
        self.mock_ocr("内嵌图片中的几何标注。")
        self.describe.side_effect = None
        self.describe.return_value = "三角形两边相等。"
        with tempfile.TemporaryDirectory() as directory:
            result = importer.extract_knowledge_document(buffer.getvalue(), "几何讲义.docx", checkpoint_dir=directory)
            self.assertEqual(self.recognize.call_count, 1, "DOCX 缩略图不能被当成教材图，也不能重复处理同图")
            self.assertEqual(self.describe.call_count, 1)
            text = "\n".join(text for _, text in result["pages"])
            self.assertIn("原生 Word 教材说明。", text)
            self.assertIn("内嵌图片中的几何标注。", text)
            self.assertIn("三角形两边相等。", text)
            self.assertEqual(text.count("内嵌图片中的几何标注。"), 1)
            self.assertTrue(all(page is None for page, _ in result["pages"]))
            self.recognize.reset_mock()
            self.describe.reset_mock()
            second = importer.extract_knowledge_document(buffer.getvalue(), "几何讲义.docx", checkpoint_dir=directory)
        self.assertEqual(result, second)
        self.recognize.assert_not_called()
        self.describe.assert_not_called()

    def test_text_only_docx_does_not_ocr_package_thumbnail(self):
        from docx import Document

        document = Document()
        document.add_paragraph("纯文字讲义，无教材图片。")
        buffer = io.BytesIO()
        document.save(buffer)
        result = importer.extract_knowledge_document(buffer.getvalue(), "纯文字.docx")
        self.assertEqual(result["pages"], [(None, "纯文字讲义，无教材图片。")])
        self.recognize.assert_not_called()
        self.describe.assert_not_called()

    def test_standalone_png_uses_ocr_visual_and_original_image_checkpoint_digest(self):
        data = self.image_bytes("Standalone equation diagram.")
        self.mock_ocr("图片中的方程 x+3=5。")
        self.describe.side_effect = None
        self.describe.return_value = "图中两条直线互相平行。"
        with tempfile.TemporaryDirectory() as directory:
            result = importer.extract_knowledge_document(data, "课堂拍照.png", checkpoint_dir=directory)
            self.recognize.assert_called_once()
            self.describe.assert_called_once()
            self.assertIn("图片中的方程 x+3=5。", self.describe.call_args.args[1])
            self.assertEqual((result["page_count"], result["ocr_pages"], result["visual_pages"]), (1, 1, 1))
            self.assertIn("图片中的方程 x+3=5。", result["pages"][0][1])
            self.assertIn("图中两条直线互相平行。", result["pages"][0][1])
            self.assertTrue((Path(directory) / hashlib.sha256(data).hexdigest() / "page-1.json").is_file())
            self.recognize.reset_mock()
            self.describe.reset_mock()
            second = importer.extract_knowledge_document(data, "课堂拍照.png", checkpoint_dir=directory)
        self.assertEqual(result, second)
        self.recognize.assert_not_called()
        self.describe.assert_not_called()

    def test_blank_pages_skip_ocr_and_keep_original_page_numbers(self):
        result = importer.extract_knowledge_pdf(self.pdf_bytes("blank", "native", "blank"))
        self.assertEqual((result["page_count"], result["indexed_pages"]), (3, 1))
        self.assertEqual([number for number, _ in result["pages"]], [2])
        self.assertEqual(len(result["warnings"]), 2)
        self.recognize.assert_not_called()
        self.describe.assert_not_called()
        with self.assertRaisesRegex(ValueError, "没有识别到可检索内容"):
            importer.extract_knowledge_pdf(self.pdf_bytes("blank"))

    def test_encrypted_or_invalid_pdf_rejected_before_ocr(self):
        for data, message in ((self.pdf_bytes("native", encrypted=True), "密码保护"),
                              (b"broken document", "无法打开 PDF"), (b"", "不能为空")):
            with self.subTest(message=message), self.assertRaisesRegex(ValueError, message):
                importer.extract_knowledge_pdf(data)
        self.recognize.assert_not_called()
        self.describe.assert_not_called()


class TesseractProtocolTests(unittest.TestCase):
    def test_native_executable_path_is_passed_as_one_argument(self):
        completed = SimpleNamespace(returncode=0, stdout=b"", stderr=b"")
        path = "/remedi/runtime/ocr build/bin/tesseract"
        with patch.object(ocr.settings, "RAG_OCR_EXECUTABLE", path), patch.object(
            ocr.subprocess, "run", return_value=completed,
        ) as run:
            ocr.recognize_text(b"image fixture")
        self.assertEqual(run.call_args.args[0][:3], [path, "stdin", "stdout"])
        self.assertFalse(run.call_args.kwargs.get("shell", False))

    def test_tsv_reconstructs_chinese_lines_english_spaces_and_confidence(self):
        header = "level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n"
        rows = [
            ("1", "1", "0", "0", "0", "0", "0", "0", "240", "300", "-1", ""),
            ("5", "1", "1", "1", "1", "1", "10", "10", "20", "20", "96.5", "数"),
            ("5", "1", "1", "1", "1", "2", "30", "10", "20", "20", "85.5", "学"),
            ("5", "1", "1", "1", "2", "1", "10", "40", "40", "20", "90", "linear"),
            ("5", "1", "1", "1", "2", "2", "50", "40", "40", "20", "88", "equation"),
        ]
        payload = (header + "\n".join("\t".join(row) for row in rows) + "\n").encode("utf-8")
        completed = SimpleNamespace(returncode=0, stdout=payload, stderr=b"")
        with patch.object(ocr.subprocess, "run", return_value=completed) as run:
            result = ocr.recognize_text(b"image fixture")
        self.assertEqual(result["text"], "数学\nlinear equation")
        self.assertAlmostEqual(result["confidence"], (0.965 + 0.855 + 0.9 + 0.88) / 4)
        command = run.call_args.args[0]
        self.assertEqual(command[:3], ["tesseract", "stdin", "stdout"])
        self.assertEqual(command[-1], "tsv")
        self.assertEqual(run.call_args.kwargs["input"], b"image fixture")
        self.assertGreater(run.call_args.kwargs["timeout"], 0)

    def test_ocr_failures_are_safe_and_never_fallback_to_visual_transcription(self):
        cases = (
            FileNotFoundError("private executable path"),
            subprocess.TimeoutExpired("private executable path", 1),
            SimpleNamespace(returncode=1, stdout=b"", stderr=b"private input contents"),
        )
        for response in cases:
            options = {"side_effect": response} if isinstance(response, Exception) else {"return_value": response}
            with self.subTest(response=response), patch.object(ocr.subprocess, "run", **options), \
                    patch.object(ocr, "describe_visual") as visual, self.assertRaises(RuntimeError) as raised:
                ocr.recognize_text(b"input")
            self.assertNotIn("private", str(raised.exception))
            visual.assert_not_called()


if __name__ == "__main__":
    unittest.main()
