"""保留文字层，扫描文字使用专用 OCR，图表按需视觉补全并逐页保存检查点。"""

import csv
import io
import json
import logging
import hashlib
import os
import re
import xml.etree.ElementTree as ET
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait

from ..config import settings
from .knowledge_ocr import describe_visual, needs_visual_context, recognize_text

logger = logging.getLogger(__name__)

SUPPORTED_KNOWLEDGE_FORMATS = (
    ".pdf", ".docx", ".pptx", ".xlsx", ".xlsm",
    ".txt", ".md", ".csv", ".tsv", ".html", ".htm", ".json", ".xml", ".yaml", ".yml",
    ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff",
)
_IMAGE_FORMATS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
_PIPELINE_VERSION = 1
_OCR_STRATEGY_VERSION = 2


def _looks_like_unusable_pdf_text(text: str) -> bool:
    """沿用 RAG 项目的质量判断，避免把明显乱码直接放入检索库。"""
    import re

    value = str(text or "").strip()
    if len(value) < 40:
        return False
    mojibake_chars = sum(1 for char in value if "\u00c0" <= char <= "\u00ff" or char == "�")
    if mojibake_chars / len(value) >= 0.03:
        return True
    tokens = [token for token in re.split(r"\s+", value) if token]
    if len(tokens) >= 30:
        single_char_tokens = sum(1 for token in tokens if len(token) == 1)
        average_token_length = sum(len(token) for token in tokens) / len(tokens)
        cjk_chars = sum(1 for char in value if "\u4e00" <= char <= "\u9fff")
        cjk_single_tokens = sum(
            1 for token in tokens
            if len(token) == 1 and "\u4e00" <= token <= "\u9fff"
        )
        if average_token_length <= 1.8 and single_char_tokens / len(tokens) >= 0.62:
            return True
        if cjk_chars >= 20 and cjk_single_tokens / cjk_chars >= 0.65:
            return True
    meaningful = sum(
        1 for char in value
        if "\u4e00" <= char <= "\u9fff" or char.isascii() and char.isalnum()
    )
    replacement = value.count("�") + value.count("?")
    return meaningful / len(value) < 0.35 or replacement / len(value) >= 0.08

def _checkpoint_path(directory, digest: str, page: int):
    if directory is None:
        return None
    root = Path(directory) / digest
    root.mkdir(parents=True, exist_ok=True)
    return root / f"page-{page}.json"


def _read_checkpoint(path):
    if path is None or not path.is_file():
        return {}
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
        return value if value.get("version") == _PIPELINE_VERSION else {}
    except (OSError, ValueError, TypeError, AttributeError):
        return {}


def _write_checkpoint(path, record):
    if path is None:
        return
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps({"version": _PIPELINE_VERSION, **record}, ensure_ascii=False), encoding="utf-8")
    os.replace(temporary, path)


def _render(page, clip=None, *, visual=False):
    try:
        import pymupdf
    except ModuleNotFoundError as exc:
        if exc.name != "pymupdf":
            raise
        import fitz as pymupdf

    rect = clip or page.rect
    scale = min(1280 / max(rect.width, rect.height), 2.0) if visual else settings.RAG_OCR_DPI / 72
    pixmap = page.get_pixmap(
        matrix=pymupdf.Matrix(scale, scale), clip=clip, alpha=False,
        colorspace=pymupdf.csRGB if visual else pymupdf.csGRAY,
    )
    return pixmap.tobytes("jpeg", jpg_quality=75) if visual else pixmap.tobytes("png")


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _filter_regions(page, candidates):
    try:
        import pymupdf
    except ModuleNotFoundError as exc:
        if exc.name != "pymupdf":
            raise
        import fitz as pymupdf

    rects = []
    area = page.rect.get_area() or 1
    for bounds in candidates:
        rect = pymupdf.Rect(bounds) & page.rect
        if rect.width < 40 or rect.height < 30 or not 0.012 <= rect.get_area() / area <= 0.82:
            continue
        if any((rect & previous).get_area() >= 0.85 * min(rect.get_area(), previous.get_area()) for previous in rects):
            continue
        rects.append(rect)
    return rects


def _raster_regions(page, image_info=None):
    """仅定位嵌入位图，供文字 OCR 使用；原生矢量绘图不重复 OCR。"""
    image_info = page.get_image_info() if image_info is None else image_info
    return _filter_regions(page, [item["bbox"] for item in image_info])


def _visual_regions(page, image_info=None):
    """定位位图与矢量图形，供按需视觉理解使用。"""
    image_info = page.get_image_info() if image_info is None else image_info
    candidates = [item["bbox"] for item in image_info]
    try:
        candidates.extend(page.cluster_drawings())
    except (AttributeError, ValueError, RuntimeError):
        pass
    return _filter_regions(page, candidates)


def _process_page(native_text, ocr_images, visual_images, context_image, cached, checkpoint):
    # 此工作线程只运行 Tesseract 子进程/网络请求，PyMuPDF 页面始终由主线程访问。
    record = dict(cached)
    warnings = []
    if not record.get("ocr_done"):
        parts = [native_text] if native_text else []
        confidence = 1.0
        ocr_parts = record.setdefault("ocr_parts", {})
        for image in ocr_images:
            image_key = hashlib.sha256(image).hexdigest()
            try:
                recognized = ocr_parts.get(image_key)
                if recognized is None:
                    recognized = recognize_text(image)
                    ocr_parts[image_key] = recognized
                    _write_checkpoint(checkpoint, record)
            except RuntimeError:
                # 即使存在原生页眉，也不能把失败的扫描正文标记为完成。
                raise RuntimeError("本页部分图片文字识别失败，已识别内容已保存，可重试继续。") from None
            if recognized["text"]:
                parts.append(recognized["text"])
                confidence = min(confidence, recognized["confidence"])
        text = "\n\n".join(parts).strip()
        if text and ocr_images and confidence < 0.55:
            warnings.append("文字识别置信度较低，请对照原件核对公式和符号。")
        record.update(text=text, confidence=confidence, ocr_done=True, ocr=bool(ocr_images),
                      ocr_strategy=_OCR_STRATEGY_VERSION, warnings=warnings)
        _write_checkpoint(checkpoint, record)
    text = record.get("text", "")
    if settings.RAG_VISUAL_ENABLED and not record.get("visual_done"):
        record["warnings"] = [item for item in record.get("warnings", []) if not item.startswith("图形解析")]
        images = visual_images
        if not images and context_image and needs_visual_context(text, record.get("confidence", 1)):
            images = [context_image]
        descriptions = []
        visual_failed = False
        visual_parts = record.setdefault("visual_parts", {})
        for image in images:
            image_key = hashlib.sha256(image).hexdigest()
            try:
                description = visual_parts.get(image_key)
                if description is None:
                    description = describe_visual(image, text)
                    visual_parts[image_key] = description
                    _write_checkpoint(checkpoint, record)
                if description:
                    descriptions.append(description)
            except RuntimeError:
                visual_failed = True
                warnings.append("图形解析暂不可用，已保留文字，可对照原件查看。")
        record["visual_text"] = "\n\n".join(descriptions)
        record["visual_done"] = not visual_failed
        record["warnings"] = list(dict.fromkeys(record.get("warnings", []) + warnings))
    else:
        record.setdefault("visual_text", "")
        record.setdefault("visual_done", not settings.RAG_VISUAL_ENABLED)
    record["complete"] = record.get("ocr_done") and record.get("visual_done")
    _write_checkpoint(checkpoint, record)
    return record


def extract_knowledge_pdf(data: bytes | os.PathLike[str] | str, *, progress=None, checkpoint_dir=None, fingerprint=None, force_visual=False) -> dict:
    try:
        import pymupdf
    except ModuleNotFoundError as exc:
        if exc.name != "pymupdf":
            raise
        import fitz as pymupdf

    if isinstance(data, bytes) and not data:
        raise ValueError("PDF 不能为空。")
    try:
        if isinstance(data, (str, os.PathLike)):
            source_path = Path(data)
            document = pymupdf.open(filename=str(source_path))
        else:
            source_path = None
            document = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ValueError("无法打开 PDF，请确认文件完整且格式正确。") from exc
    with document:
        if document.needs_pass:
            raise ValueError("请先解除 PDF 密码保护再导入。")
        if document.page_count < 1:
            raise ValueError("PDF 没有可读取的页面。")
        results, failures = {}, {}
        total = document.page_count
        completed_text_pages = 0
        last_report_stage = "extracting"
        digest = fingerprint or (
            hashlib.sha256(data).hexdigest() if source_path is None else _sha256_file(source_path)
        )

        def report(stage=None):
            nonlocal last_report_stage
            if stage is not None:
                last_report_stage = stage
            if progress:
                # 由完成事件增量维护，避免每页重新扫描此前所有结果。
                progress({"stage": last_report_stage, "completed_pages": completed_text_pages,
                          "total_pages": total, "message": f"已处理 {completed_text_pages}/{total} 页"})

        def collect(pending):
            nonlocal completed_text_pages
            completed, _ = wait(pending, return_when=FIRST_COMPLETED)
            for future in completed:
                number = pending.pop(future)
                try:
                    record = future.result()
                    results[number] = record
                    completed_text_pages += int(bool(record.get("ocr_done")))
                except RuntimeError as exc:
                    failures[number] = str(exc)
                    logger.warning("教材第 %s 页文字识别失败（%s）", number, type(exc).__name__)
                report()

        with ThreadPoolExecutor(max_workers=settings.RAG_OCR_WORKERS, thread_name_prefix="knowledge-page") as pool:
            pending = {}
            for index, page in enumerate(document):
                number = index + 1
                checkpoint = _checkpoint_path(checkpoint_dir, digest, number)
                cached = _read_checkpoint(checkpoint)
                if cached.get("ocr_strategy") != _OCR_STRATEGY_VERSION:
                    # 旧检查点可能把矢量图形也送去 OCR；仅重算文本步骤并复用图像 OCR 结果。
                    cached = dict(cached)
                    cached["ocr_done"] = False
                    cached["complete"] = False
                if cached.get("complete"):
                    results[number] = cached
                    completed_text_pages += int(bool(cached.get("ocr_done")))
                    report()
                    continue
                text = page.get_text("text", sort=True).replace("\x00", "").strip()
                scan = not text or _looks_like_unusable_pdf_text(text)
                image_info = page.get_image_info()
                # 带原生页眉的扫描正文仍需 OCR，不能因为有少量文字层而漏掉正文。
                full_scan_image = any(
                    (pymupdf.Rect(item["bbox"]) & page.rect).get_area() / max(page.rect.get_area(), 1) > 0.82
                    for item in image_info
                )
                # 无任何绘制内容的空白页直接跳过，不加载 OCR。
                blank = not text and not image_info and not page.get_drawings()
                raster_regions = [] if scan else _raster_regions(page, image_info)
                visual_regions = (
                    _visual_regions(page, image_info)
                    if not scan and settings.RAG_VISUAL_ENABLED else []
                )
                ocr_images = [] if cached.get("ocr_done") or blank else (
                    [_render(page)] if scan or full_scan_image else [_render(page, rect) for rect in raster_regions]
                )
                visuals = [_render(page, rect, visual=True) for rect in visual_regions]
                needs_page_context = force_visual or scan or (
                    not visual_regions and needs_visual_context(text, cached.get("confidence", 1))
                )
                context_image = (
                    _render(page, visual=True)
                    if settings.RAG_VISUAL_ENABLED and not blank and needs_page_context else None
                )
                if force_visual and context_image:
                    visuals = [context_image]
                report("ocr" if ocr_images else "visual" if visuals else "extracting")
                future = pool.submit(_process_page, "" if scan else text, ocr_images, visuals,
                                     context_image, cached, checkpoint)
                pending[future] = number
                if len(pending) >= settings.RAG_OCR_WORKERS * 2:
                    collect(pending)
            while pending:
                collect(pending)
        pages, warnings, ocr_pages, visual_pages = [], [], 0, 0
        for number in sorted(results):
            record = results[number]
            text = record.get("text", "")
            visual = record.get("visual_text", "")
            if visual:
                text += "\n\n[图形语义补全]\n" + visual
                visual_pages += 1
            if text.strip():
                pages.append((number, text.strip()))
            else:
                warnings.append(f"第 {number} 页未识别到文字或图形内容。")
            ocr_pages += bool(record.get("ocr"))
            warnings.extend(f"第 {number} 页：{warning}" for warning in record.get("warnings", []))
        if failures:
            # 缺页不能静默标记为整本完成；成功页面已缓存，重试只处理失败页。
            raise RuntimeError(f"第 {'、'.join(map(str, sorted(failures)))} 页文字识别失败；已完成页面已保存，重试可继续。")
        # 图形理解是文字 OCR 的补充。视觉模型暂时超时不能阻止已经识别的
        # 正文入库；未完成的图形页仍保留不完整检查点，后续重试可复用 OCR。
        if not pages:
            raise ValueError("文件中没有识别到可检索内容，请检查扫描清晰度。")
        return {"pages": pages, "page_count": total, "indexed_pages": len(pages),
                "ocr_pages": ocr_pages, "visual_pages": visual_pages, "warnings": warnings}


def _decode_text(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16").replace("\x00", "").strip()
        except UnicodeDecodeError:
            pass
    for encoding in ("utf-8-sig", "gb18030", "cp1252"):
        try:
            return data.decode(encoding).replace("\x00", "").strip()
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文件编码，请另存为 UTF-8 后再导入。")


class _VisibleHtmlText(HTMLParser):
    _BLOCK_TAGS = {"address", "article", "br", "div", "footer", "h1", "h2", "h3", "h4",
                   "header", "li", "main", "ol", "p", "section", "table", "td", "th", "tr", "ul"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden_depth = 0

    def handle_starttag(self, tag, _attrs):
        if tag in {"script", "style", "noscript"}:
            self.hidden_depth += 1
        elif not self.hidden_depth and tag in self._BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript"} and self.hidden_depth:
            self.hidden_depth -= 1
        elif not self.hidden_depth and tag in self._BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden_depth:
            self.parts.append(data)

    def text(self) -> str:
        return re.sub(r"[ \t]+", " ", "".join(self.parts)).strip()


def _office_source(data: bytes | os.PathLike[str] | str):
    return str(data) if isinstance(data, (str, os.PathLike)) else io.BytesIO(data)


def _docx_text(data: bytes | os.PathLike[str] | str) -> str:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(_office_source(data))
    parts = []
    for block in document.element.body.iterchildren():
        if block.tag == qn("w:p"):
            text = Paragraph(block, document).text.strip()
            if text:
                parts.append(text)
        elif block.tag == qn("w:tbl"):
            table = Table(block, document)
            for row in table.rows:
                cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
    return "\n\n".join(parts).strip()


def _pptx_pages(data: bytes | os.PathLike[str] | str) -> tuple[list[tuple[None, str]], list[str]]:
    from pptx import Presentation

    presentation = Presentation(_office_source(data))
    pages = []
    skipped = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
        text = "\n".join(parts).strip()
        if text:
            pages.append((None, f"幻灯片 {index}\n{text}"))
        else:
            skipped.append(str(index))
    warnings = ["以下幻灯片未提取到文字：" + "、".join(skipped) + "。"] if skipped else []
    return pages, warnings


def _xlsx_text(data: bytes | os.PathLike[str] | str) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(_office_source(data), read_only=True, data_only=False)
    parts = []
    try:
        for sheet in workbook.worksheets:
            parts.append(f"工作表：{sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value).strip() for value in row]
                if any(values):
                    parts.append(" | ".join(values))
    finally:
        workbook.close()
    return "\n".join(parts).strip()


def _office_images(data: bytes | os.PathLike[str] | str, extension: str):
    """复用现有 Office 解析库，图片型 Word/课件同样进入 OCR/图形链路。"""
    if extension == ".docx":
        from docx import Document
        from docx.parts.image import ImagePart

        for part in Document(_office_source(data)).part.package.parts:
            if isinstance(part, ImagePart):
                yield part.blob
    elif extension == ".pptx":
        from pptx import Presentation

        def pictures(shapes):
            for shape in shapes:
                if hasattr(shape, "image"):
                    yield shape.image.blob
                elif hasattr(shape, "shapes"):
                    yield from pictures(shape.shapes)

        for slide in Presentation(_office_source(data)).slides:
            yield from pictures(slide.shapes)


def _extract_picture(data: bytes | os.PathLike[str] | str, *, progress=None, checkpoint_dir=None, fingerprint=None):
    try:
        import pymupdf
    except ModuleNotFoundError as exc:
        if exc.name != "pymupdf":
            raise
        import fitz as pymupdf

    try:
        if isinstance(data, (str, os.PathLike)):
            source_path = Path(data)
            image = pymupdf.open(filename=str(source_path))
            digest = fingerprint or _sha256_file(source_path)
        else:
            image = pymupdf.open(stream=data)
            digest = fingerprint or hashlib.sha256(data).hexdigest()
        with image:
            pdf = image.convert_to_pdf()
    except Exception:
        raise ValueError("无法读取图片，请检查文件是否完整或转换为 PNG/JPEG。") from None
    return extract_knowledge_pdf(pdf, progress=progress, checkpoint_dir=checkpoint_dir,
                                 fingerprint=digest, force_visual=True)


def _extract_non_pdf(data: bytes | os.PathLike[str] | str, extension: str, *, progress=None, checkpoint_dir=None) -> dict:
    warnings = []
    if extension == ".docx":
        try:
            text = _docx_text(data)
        except Exception as exc:
            raise ValueError("无法读取 Word 文档，请确认这是未损坏的 DOCX 文件。") from exc
        pages = [(None, text)] if text else []
    elif extension == ".pptx":
        try:
            pages, warnings = _pptx_pages(data)
        except Exception as exc:
            raise ValueError("无法读取 PowerPoint 文档，请确认这是未损坏的 PPTX 文件。") from exc
        text = "\n\n".join(page_text for _, page_text in pages)
    elif extension in {".xlsx", ".xlsm"}:
        try:
            text = _xlsx_text(data)
        except Exception as exc:
            raise ValueError("无法读取 Excel 文档，请确认这是未损坏的 XLSX/XLSM 文件。") from exc
        pages = [(None, text)] if text else []
    else:
        if isinstance(data, (str, os.PathLike)):
            data = Path(data).read_bytes()
        text = _decode_text(data)
        if extension in {".csv", ".tsv"}:
            delimiter = "\t" if extension == ".tsv" else ","
            try:
                rows = csv.reader(io.StringIO(text), delimiter=delimiter)
                text = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows).strip()
            except csv.Error as exc:
                raise ValueError("无法读取表格文本，请检查 CSV/TSV 文件格式。") from exc
        elif extension in {".html", ".htm"}:
            parser = _VisibleHtmlText()
            parser.feed(text)
            text = parser.text()
        elif extension == ".json":
            try:
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            except json.JSONDecodeError as exc:
                raise ValueError("无法读取 JSON 文件，请检查文件格式。") from exc
        elif extension == ".xml":
            try:
                text = "\n".join(value.strip() for value in ET.fromstring(text).itertext() if value.strip())
            except ET.ParseError as exc:
                raise ValueError("无法读取 XML 文件，请检查文件格式。") from exc
        pages = [(None, text)] if text else []

    if extension in {".docx", ".pptx"}:
        seen = set()
        for image_index, image in enumerate(_office_images(data, extension), 1):
            digest = hashlib.sha256(image).hexdigest()
            if digest in seen:
                continue
            seen.add(digest)
            if progress:
                progress({"stage": "visual", "message": f"正在处理文档图片 {image_index}"})
            try:
                extracted = _extract_picture(image, checkpoint_dir=checkpoint_dir)
            except ValueError:
                warnings.append(f"文档图片 {image_index} 格式无法读取，请对照原件核对。")
                continue
            pages.append((None, f"文档图片 {image_index}\n" + "\n\n".join(t for _, t in extracted["pages"])))
            warnings.extend(extracted["warnings"])
        text = "\n\n".join(t for _, t in pages)
    if not text.strip():
        raise ValueError("文件中没有提取到可检索的文字内容。")
    return {"pages": pages, "warnings": warnings}


def extract_knowledge_document(data: bytes | os.PathLike[str] | str, filename: str, *, progress=None, checkpoint_dir=None, fingerprint=None) -> dict:
    """按文件扩展名读取常见教学资料，统一交给学科 RAG 切分和索引。"""
    if not data:
        raise ValueError("资料文件不能为空。")
    extension = PurePosixPath((filename or "").replace("\\", "/")).suffix.lower()
    if extension not in SUPPORTED_KNOWLEDGE_FORMATS:
        supported = "、".join(SUPPORTED_KNOWLEDGE_FORMATS)
        raise ValueError(f"暂不支持此文件格式。支持：{supported}。")
    if extension == ".pdf":
        return extract_knowledge_pdf(
            data, progress=progress, checkpoint_dir=checkpoint_dir, fingerprint=fingerprint,
        )
    if extension in _IMAGE_FORMATS:
        return _extract_picture(
            data, progress=progress, checkpoint_dir=checkpoint_dir, fingerprint=fingerprint,
        )
    if progress:
        progress({"stage": "extracting", "completed_pages": 0, "total_pages": None, "message": "正在提取文档内容"})
    return _extract_non_pdf(data, extension, progress=progress, checkpoint_dir=checkpoint_dir)
